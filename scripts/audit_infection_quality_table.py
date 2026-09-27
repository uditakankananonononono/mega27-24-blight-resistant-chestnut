"""Replay published supplementary per-library read-quality metrics (not raw reads)."""
from pathlib import Path
import hashlib,json,statistics,collections,math
from docx import Document
from bs4 import BeautifulSoup
R=Path(__file__).resolve().parents[1]
p=R/'data/sources/12870_2023_4072_MOESM2_ESM.docx'
xml=R/'data/sources/pmc9901152.xml'
source=BeautifulSoup(xml.read_text(),'xml')
media=source.find('supplementary-material',id='MOESM2').find('media')
expected=next(x.split(' ',1)[1].split('?>')[0] for x in str(media).split('<?') if x.startswith('suppdata-md5 '))
assert hashlib.md5(p.read_bytes()).hexdigest()==expected
D=Document(p);assert len(D.tables)==1
t=D.tables[0];assert len(t.rows)==16 and len(t.columns)==7
headers=[x.text.strip() for x in t.rows[0].cells]
assert headers==['Sample Name','Clean data','Mapping Reads','Q20','Q30','GC Content(%)','Reads aligned(%)']
rows=[]
for row in t.rows[1:]:
 c=[x.text.strip() for x in row.cells]
 label=c[0]
 group=label.rsplit('-',1)[0]
 assert group in ['Mock0h','Mock3h','Mock9h','T3h','T9h']
 clean=int(c[1]);mapped=int(c[2]);percent=float(c[6].rstrip('%'))
 rows.append({'label':label,'group':group,'clean_reads':clean,'mapped_reads':mapped,
  'q20_percent':float(c[3].rstrip('%')),'q30_percent':float(c[4].rstrip('%')),
  'gc_percent':float(c[5].rstrip('%')),'reported_mapping_percent':percent,
  'calculated_mapping_percent':100*mapped/clean})
assert len(rows)==15 and len({x['label'] for x in rows})==15
assert all(abs(x['reported_mapping_percent']-x['calculated_mapping_percent'])<.011 for x in rows)
groups={}
for k in ['Mock0h','Mock3h','Mock9h','T3h','T9h']:
 a=[x for x in rows if x['group']==k];assert len(a)==3
 groups[k]={'n':3,'mean_mapping_percent':statistics.mean(x['calculated_mapping_percent'] for x in a),
            'min_mapping_percent':min(x['calculated_mapping_percent'] for x in a),
            'max_mapping_percent':max(x['calculated_mapping_percent'] for x in a),
            'mean_clean_reads':statistics.mean(x['clean_reads'] for x in a)}
print(json.dumps({'source':'https://www.ebi.ac.uk/europepmc/webservices/rest/PMC9901152/supplementaryFiles',
 'paper':'https://pmc.ncbi.nlm.nih.gov/articles/PMC9901152/',
 'docx_md5':expected,'docx_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),
 'sample_count':len(rows),'group_summary':groups,'rows':rows,
 'total_clean_reads':sum(x['clean_reads'] for x in rows),
 'reported_vs_calculated_mapping_max_absolute_percentage_points':max(abs(x['reported_mapping_percent']-x['calculated_mapping_percent']) for x in rows),
 'limits':['Original supplementary table reports processed read quality only, not raw sequencing payload or per-gene expression.','Three replicates per group are one published cohort and 15 sample rows do not equal 15 independent studies or 15 fetched-and-used accession payloads.','Mapping differences are descriptive QC, not resistance mechanism, causal treatment effect or independent benchmark.'],
 'gate_credit':{'external_services':0,'fetched_and_used_accession_datasets':0,'audited_derivations':0,'paper_pages':0}},indent=2,sort_keys=True))
