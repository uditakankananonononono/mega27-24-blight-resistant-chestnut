"""Cross-check the article's S1 RNA-seq sample table with Figshare expression headers."""
from pathlib import Path
import hashlib,json,re
import openpyxl
R=Path(__file__).resolve().parents[1]
s=R/'data/sources/pmc12103606_supplementary_table1.xlsx'
h=R/'data/sources/curation2025_expression_headers.json'
w=openpyxl.load_workbook(s,read_only=True,data_only=True)
assert w.sheetnames==['Sheet1']
rows=list(w.active.iter_rows(values_only=True))
assert rows[0][0]=='Table S1. The sample information of RNA-Seq data.'
assert rows[2][0]=='SRR ID'
records=[{'excel_row':i,'run_accession':str(r[0]).strip(), 'tissue':r[1],
          'development_stage':r[2], 'cultivar':r[3], 'library_layout':r[4]}
         for i,r in enumerate(rows,1) if r[0] and re.fullmatch(r'[SED]RR\d+',str(r[0]).strip())]
ids=[q['run_accession'] for q in records]
assert len(ids)==len(set(ids))
x=json.loads(h.read_text())
f=set(x['gene_expression/all_exp_fpkm.txt.gz'][1:])
raw={k:set(v[1:]) for k,v in x.items() if k.startswith('gene_expression/all_exp_raw/')}
assert len(f)==213 and all(f==set(x[k][1:]) for k in ['gene_expression/all_exp_fpkm_corr.txt.gz'])
assert f==set().union(*raw.values())
supp=set(ids)
# Compare exact labels first, then only normalize suffixes where one unambiguous
# SRR number is the leading token (no guess at a sample relationship).
base=lambda label:re.match(r'^[SED]RR\d+',label).group()
normalized={base(q) for q in f}
assert len(normalized)==len(f)
assert len(supp)==212 and sorted(supp-normalized)==[] and sorted(normalized-supp)==['SRR21681158']
renamed=[{'table_id':q,'matrix_label':v} for q in sorted(supp) for v in sorted(f) if base(v)==q and v!=q]
assert renamed==[{'table_id':'SRR8383229','matrix_label':'SRR8383229_SHW_Gall_A2'}]
out={'article_url':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12103606/',
     'supplement_zip_url':'https://www.ebi.ac.uk/europepmc/webservices/rest/PMC12103606/supplementaryFiles',
     'figshare_dataset_url':'https://springernature.figshare.com/articles/dataset/Comprehensive_curation_and_validation_of_genomic_datasets_for_chestnut/27060067',
     'supplement_xlsx_sha256':hashlib.sha256(s.read_bytes()).hexdigest(),
     'figshare_archived_header_json_sha256':hashlib.sha256(h.read_bytes()).hexdigest(),
     'supplement_id_rows':len(ids), 'supplement_distinct_run_ids':len(supp),
     'figshare_fpkm_columns':len(f),'figshare_raw_columns_by_layout':{k:len(v) for k,v in raw.items()},
     'exact_table_only_labels':sorted(supp-f), 'exact_matrix_only_labels':sorted(f-supp),
     'suffix_correspondence_by_run_prefix':renamed,
     'prefix_normalized_table_only':sorted(supp-normalized),
     'prefix_normalized_matrix_only':sorted(normalized-supp),
     'extra_matrix_run_id':'SRR21681158',
     'table_records':records,
     'finding':'After one transparent SRR-prefix label correspondence, 212 unique Supplementary Table 1 run IDs are covered by the 213 Figshare matrix sample columns, with one additional matrix run ID (SRR21681158) absent from S1.',
     'limits':['The matrix column and table row counts are metadata counts, not separately downloaded RNA-seq reads or expression analyses.',
               'The source does not explain why SRR21681158 is absent from S1; do not label it an invalid biological sample or assume a processing error.',
               'The SRR8383229 suffix is a label-level correspondence only, not independent verification of the sample identity.',
               'No accession payload, phenotype, comparator, new discovery or project gate credit.'],
     'gate_credit':{'external_services':0,'fetched_and_used_accession_datasets':0,'audited_derivations':0,'paper_pages':0}}
print(json.dumps(out,indent=2,sort_keys=True))
